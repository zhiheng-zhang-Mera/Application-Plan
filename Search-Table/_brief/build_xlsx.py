#!/usr/bin/env python3
"""把汇总 CSV 生成一个多工作表 XLSX（优先表格阅读）。

用法：python Search-Table/_brief/build_xlsx.py
输入：Search-Table/data/search-table.csv
输出：Search-Table/data/search-table.xlsx
"""
from __future__ import annotations

import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "search-table.csv"
OUT = ROOT / "data" / "search-table.xlsx"

HEADERS = [
    ("region", "区域", 16),
    ("state_province", "州/省", 14),
    ("university", "学校", 34),
    ("program", "项目", 46),
    ("deadline_2027fall", "2027 Fall 截止/轮次", 40),
    ("english_waiver", "英语豁免", 46),
    ("gre", "GRE", 34),
    ("outreach", "套磁", 30),
    ("interview", "面试", 26),
    ("supervisor_acceptance", "必须导师接收", 26),
    ("funding", "资助制度", 52),
    ("other_materials", "其它提交材料", 46),
    ("academic_floor", "学术门槛", 34),
    ("match_rate", "匹配率", 16),
    ("key_risk", "关键风险", 34),
    ("sources", "来源", 60),
    ("verified_date", "核验日", 12),
]

FILL = {
    "高": PatternFill("solid", fgColor="C6EFCE"),
    "中高": PatternFill("solid", fgColor="E2EFDA"),
    "中": PatternFill("solid", fgColor="FFF2CC"),
    "低": PatternFill("solid", fgColor="FCE4D6"),
    "极低": PatternFill("solid", fgColor="F8CBAD"),
    "排除": PatternFill("solid", fgColor="D9D9D9"),
}

THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def band(match_rate: str) -> str:
    head = re.split(r"[（(]", match_rate)[0].strip()
    for key in ("未评级", "排除", "极低", "中高", "高", "中", "低"):
        if head.startswith(key):
            return key
    return ""


def load() -> list[dict[str, str]]:
    with SRC.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def style_sheet(ws, rows: list[dict[str, str]]) -> None:
    ws.append([title for _, title, _ in HEADERS])
    for idx, (_, _, width) in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    for row in rows:
        ws.append([row.get(key, "") for key, _, _ in HEADERS])
        r = ws.max_row
        tag = band(row.get("match_rate", ""))
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN
        if tag in FILL:
            match_col = [k for k, _, _ in HEADERS].index("match_rate") + 1
            ws.cell(row=r, column=match_col).fill = FILL[tag]
        ws.row_dimensions[r].height = 88

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"


def main() -> int:
    if not SRC.exists():
        print(f"缺少 {SRC}，请先运行 build_csv.py", file=sys.stderr)
        return 1

    rows = load()
    wb = Workbook()

    # 1) 总表
    ws = wb.active
    ws.title = "全部"
    style_sheet(ws, rows)

    # 2) 分区
    by_region: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        by_region[row["region"]].append(row)
    for region in sorted(by_region):
        name = region.replace("/", "-")[:28]
        style_sheet(wb.create_sheet(name), by_region[region])

    # 3) 统计
    ws = wb.create_sheet("统计")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    for col in ("E", "F", "G", "H"):
        ws.column_dimensions[col].width = 12

    ws.append(["区域", "高", "中高", "中", "低", "极低", "排除", "合计"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for region in sorted(by_region):
        counter = Counter(band(r["match_rate"]) for r in by_region[region])
        ws.append([
            region,
            counter.get("高", 0),
            counter.get("中高", 0),
            counter.get("中", 0),
            counter.get("低", 0),
            counter.get("极低", 0),
            counter.get("排除", 0),
            len(by_region[region]),
        ])
    total = Counter(band(r["match_rate"]) for r in rows)
    ws.append([
        "合计",
        total.get("高", 0),
        total.get("中高", 0),
        total.get("中", 0),
        total.get("低", 0),
        total.get("极低", 0),
        total.get("排除", 0),
        len(rows),
    ])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"合计 {len(rows)} 行 / {len(by_region)} 个分区 → {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
